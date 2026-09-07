"""Parse official ABS Recorded Crime – Offenders workbooks.

ABS counts unique alleged offenders proceeded against in a financial year.
That is a different statistical product from CSA alleged offender incidents.
"""

from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.config import ABS_STATISTICAL_PRODUCT, source_path

FINANCIAL_YEARS = [f"{year}–{str(year + 1)[-2:]}" for year in range(2008, 2025)]

JURISDICTION_ALIASES = {
    "nsw": "New South Wales",
    "new south wales": "New South Wales",
    "vic": "Victoria",
    "vic.": "Victoria",
    "victoria": "Victoria",
    "qld": "Queensland",
    "queensland": "Queensland",
    "sa": "South Australia",
    "south australia": "South Australia",
    "wa": "Western Australia",
    "western australia": "Western Australia",
    "tas": "Tasmania",
    "tas.": "Tasmania",
    "tasmania": "Tasmania",
    "nt": "Northern Territory",
    "northern territory": "Northern Territory",
    "act": "Australian Capital Territory",
    "australian capital territory": "Australian Capital Territory",
    "australia": "Australia",
}

SEX_ALIASES = {
    "males": "Males",
    "male": "Males",
    "females": "Females",
    "female": "Females",
    "persons": "Persons",
}


def strip_footnotes(value: object) -> str:
    """Remove ABS footnote markers such as (d), (h)(i) and the registered mark."""
    text = str(value or "").replace("\n", " ").replace("®", "")
    text = re.sub(r"\([a-z]{1,3}\)", "", text, flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", text).strip(" .")


def normalise_year(value: object) -> str | None:
    """Normalise ABS financial-year labels to an en-dash form (2024–25)."""
    if value is None:
        return None
    text = str(value).strip().replace("—", "–").replace("-", "–")
    match = re.search(r"(20\d{2})–(\d{2})", text)
    if not match:
        return None
    return f"{match.group(1)}–{match.group(2)}"


def year_end(financial_year: str) -> int:
    """Return the calendar year in which the financial year ends (2024–25 → 2025)."""
    normalised = normalise_year(financial_year)
    if not normalised:
        raise ValueError(f"Unrecognised financial year: {financial_year}")
    return 2000 + int(normalised.split("–")[1])


def normalise_jurisdiction(value: object) -> str | None:
    key = strip_footnotes(value).lower().rstrip(".")
    return JURISDICTION_ALIASES.get(key)


def normalise_sex(value: object) -> str | None:
    key = strip_footnotes(value).lower()
    return SEX_ALIASES.get(key)


def offence_level(label: object) -> str | None:
    """Classify an ABS principal-offence row as total, division, subdivision or other."""
    text = strip_footnotes(label)
    if not text:
        return None
    if text.lower().startswith("total"):
        return "total"
    if text.lower().startswith("fare evasion"):
        return "other"
    match = re.match(r"^(\d{2,3})\b", text)
    if not match:
        return None
    return "division" if len(match.group(1)) == 2 else "subdivision"


def sheet_rows(path: Path, sheet: str, max_col: int = 40) -> list[tuple]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(workbook[sheet].iter_rows(max_col=max_col, values_only=True))
    finally:
        workbook.close()


def _numeric(value: object) -> float | None:
    if value is None or value == "" or value == "np":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _year_header_indexes(header_row: tuple) -> list[int]:
    """Return column indexes of the first 17 financial-year headers (counts, not rates)."""
    indexes: list[int] = []
    seen: set[str] = set()
    for idx, cell in enumerate(header_row):
        year = normalise_year(cell)
        if year and year not in seen:
            indexes.append(idx)
            seen.add(year)
        if len(indexes) == len(FINANCIAL_YEARS):
            break
    return indexes


def _series_from_wide_row(
    header_row: tuple,
    count_row: tuple,
    *,
    source_workbook: str,
    source_table: str,
    jurisdiction: str,
    sex: str,
    age_group: str,
    principal_offence: str,
    offence_level_name: str,
) -> list[dict]:
    year_indexes = _year_header_indexes(header_row)
    if len(year_indexes) != len(FINANCIAL_YEARS):
        raise ValueError(f"{source_table} does not contain 17 financial-year headers")
    rate_offset = year_indexes[-1] - year_indexes[0] + 1
    records = []
    for position, year_idx in enumerate(year_indexes):
        year = FINANCIAL_YEARS[position]
        count = _numeric(count_row[year_idx] if year_idx < len(count_row) else None)
        rate_idx = year_idx + rate_offset
        rate = _numeric(count_row[rate_idx] if rate_idx < len(count_row) else None)
        if count is None and rate is None:
            continue
        records.append(
            _record(
                source_workbook=source_workbook,
                source_table=source_table,
                financial_year=year,
                jurisdiction=jurisdiction,
                sex=sex,
                age_group=age_group,
                principal_offence=principal_offence,
                offence_level_name=offence_level_name,
                offender_count=count,
                rate=rate,
            )
        )
    return records


def _record(
    *,
    source_workbook: str,
    source_table: str,
    financial_year: str,
    jurisdiction: str,
    sex: str,
    age_group: str,
    principal_offence: str,
    offence_level_name: str,
    offender_count: float | None,
    rate: float | None,
) -> dict:
    return {
        "source_workbook": source_workbook,
        "source_table": source_table,
        "financial_year": financial_year,
        "year_end": year_end(financial_year),
        "jurisdiction": jurisdiction,
        "sex": sex,
        "age_group": age_group,
        "principal_offence": principal_offence,
        "offence_level": offence_level_name,
        "offender_count": offender_count,
        "rate_per_100000_age_10_plus": rate,
        "statistical_product": "ABS unique alleged offenders proceeded against",
        "caveat": ABS_STATISTICAL_PRODUCT,
    }


def parse_table6_states_latest(path: Path) -> pd.DataFrame:
    """Table 6: offenders by principal offence and jurisdiction, 2024–25 only."""
    rows = sheet_rows(path, "Table 6")
    header = rows[5]
    jurisdictions: list[tuple[int, str]] = []
    seen: set[str] = set()
    for idx, cell in enumerate(header):
        name = normalise_jurisdiction(cell)
        if name and name not in seen:
            jurisdictions.append((idx, name))
            seen.add(name)
        if len(jurisdictions) == 8:
            break
    if len(jurisdictions) != 8:
        raise ValueError("Table 6 did not expose eight state and territory columns")
    rate_offset = jurisdictions[-1][0] - jurisdictions[0][0] + 1
    records: list[dict] = []
    for row in rows[7:]:
        level = offence_level(row[0])
        if level is None:
            continue
        label = strip_footnotes(row[0])
        if level == "total":
            label = "Total"
        for col_idx, jurisdiction in jurisdictions:
            records.append(
                _record(
                    source_workbook=path.name,
                    source_table="Table 6",
                    financial_year="2024–25",
                    jurisdiction=jurisdiction,
                    sex="Persons",
                    age_group="All ages 10+",
                    principal_offence=label,
                    offence_level_name=level,
                    offender_count=_numeric(row[col_idx] if col_idx < len(row) else None),
                    rate=_numeric(
                        row[col_idx + rate_offset]
                        if col_idx + rate_offset < len(row)
                        else None
                    ),
                )
            )
        if level == "total":
            break
    return pd.DataFrame(records)


def parse_table8_victoria_sex(path: Path) -> pd.DataFrame:
    """Table 8: Victoria offenders by sex, 2008–09 to 2024–25."""
    rows = sheet_rows(path, "Table 8")
    header = rows[5]
    current_sex: str | None = None
    records: list[dict] = []
    for row in rows[6:]:
        maybe_sex = normalise_sex(row[1]) if row[1] and not row[0] else None
        if maybe_sex:
            current_sex = maybe_sex
            continue
        if offence_level(row[0]) != "total" or current_sex is None:
            continue
        records.extend(
            _series_from_wide_row(
                header,
                row,
                source_workbook=path.name,
                source_table="Table 8",
                jurisdiction="Victoria",
                sex=current_sex,
                age_group="All ages 10+",
                principal_offence="Total",
                offence_level_name="total",
            )
        )
    return pd.DataFrame(records)


def parse_table1_australia(path: Path) -> pd.DataFrame:
    """Table 1: Australia offender totals, 2008–09 to 2024–25."""
    rows = sheet_rows(path, "Table 1")
    header = rows[5]
    for row in rows[6:]:
        if offence_level(row[0]) == "total":
            return pd.DataFrame(
                _series_from_wide_row(
                    header,
                    row,
                    source_workbook=path.name,
                    source_table="Table 1",
                    jurisdiction="Australia",
                    sex="Persons",
                    age_group="All ages 10+",
                    principal_offence="Total",
                    offence_level_name="total",
                )
            )
    raise ValueError("Table 1 total row was not found")


def parse_table1_australia_divisions(path: Path) -> pd.DataFrame:
    """Table 1 ANZSOC divisions for the latest financial year only."""
    rows = sheet_rows(path, "Table 1")
    header = rows[5]
    year_indexes = _year_header_indexes(header)
    latest_idx = year_indexes[-1]
    rate_idx = latest_idx + (year_indexes[-1] - year_indexes[0] + 1)
    records: list[dict] = []
    for row in rows[6:]:
        level = offence_level(row[0])
        if level not in {"division", "total"}:
            continue
        label = "Total" if level == "total" else strip_footnotes(row[0])
        records.append(
            _record(
                source_workbook=path.name,
                source_table="Table 1",
                financial_year="2024–25",
                jurisdiction="Australia",
                sex="Persons",
                age_group="All ages 10+",
                principal_offence=label,
                offence_level_name=level,
                offender_count=_numeric(row[latest_idx] if latest_idx < len(row) else None),
                rate=_numeric(row[rate_idx] if rate_idx < len(row) else None),
            )
        )
        if level == "total":
            break
    return pd.DataFrame(records)


def parse_table20_youth(path: Path) -> pd.DataFrame:
    """Table 20: youth offenders (10–17) by state, 2008–09 to 2024–25."""
    rows = sheet_rows(path, "Table 20")
    header = rows[5]
    current: str | None = None
    records: list[dict] = []
    for row in rows[6:]:
        section = normalise_jurisdiction(row[1]) if row[1] and not row[0] else None
        if section:
            current = section
            continue
        label = strip_footnotes(row[0])
        if current is None or not label.lower().startswith("youth offender"):
            continue
        records.extend(
            _series_from_wide_row(
                header,
                row,
                source_workbook=path.name,
                source_table="Table 20",
                jurisdiction=current,
                sex="Persons",
                age_group="Youth 10-17",
                principal_offence="Total",
                offence_level_name="total",
            )
        )
    return pd.DataFrame(records)


def load_abs_tables(
    australia_path: Path | None = None,
    states_path: Path | None = None,
    youth_path: Path | None = None,
) -> dict[str, pd.DataFrame]:
    """Return tidy ABS tables ready for SQLite and CSV export."""
    australia = australia_path or source_path("abs_offenders_australia")
    states = states_path or source_path("abs_offenders_states")
    youth = youth_path or source_path("abs_youth_offenders")

    states_latest = parse_table6_states_latest(states)
    victoria_trend = parse_table8_victoria_sex(states)
    australia_trend = parse_table1_australia(australia)
    australia_divisions = parse_table1_australia_divisions(australia)
    youth_trend = parse_table20_youth(youth)

    trend = pd.concat(
        [
            victoria_trend,
            australia_trend,
        ],
        ignore_index=True,
    )
    latest_totals = states_latest.loc[states_latest["offence_level"].eq("total")].copy()
    australia_total = australia_trend.loc[australia_trend["financial_year"].eq("2024–25")].copy()
    states_and_australia = pd.concat([latest_totals, australia_total], ignore_index=True)

    combined = pd.concat(
        [
            states_latest,
            victoria_trend,
            australia_trend,
            australia_divisions,
            youth_trend,
        ],
        ignore_index=True,
    )
    combined = combined.drop_duplicates(
        subset=[
            "source_table",
            "financial_year",
            "jurisdiction",
            "sex",
            "age_group",
            "principal_offence",
        ]
    )
    return {
        "abs_offenders": combined,
        "abs_offenders_states_latest": states_latest,
        "abs_offenders_trend": trend,
        "abs_offenders_states_totals": states_and_australia,
        "abs_youth_offenders": youth_trend,
    }
